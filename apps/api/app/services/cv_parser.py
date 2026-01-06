"""CV/Resume parsing service."""

import io
import re
from typing import Any, Optional

import pdfplumber
from docx import Document

from app.services.llm import llm_provider


def extract_text_from_pdf(content: bytes) -> str:
    """Extract text from PDF file."""
    text_parts = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n\n".join(text_parts)


def extract_text_from_docx(content: bytes) -> str:
    """Extract text from DOCX file."""
    doc = Document(io.BytesIO(content))
    text_parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # Also extract from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)

    return "\n".join(text_parts)


def extract_text_from_file(content: bytes, file_type: str) -> str:
    """Extract text from file based on type."""
    if file_type.lower() in ("pdf", "application/pdf"):
        return extract_text_from_pdf(content)
    elif file_type.lower() in (
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ):
        return extract_text_from_docx(content)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def mask_pii(text: str) -> str:
    """Mask PII (email, phone) in text."""
    # Mask emails
    email_pattern = r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b"
    text = re.sub(email_pattern, "[EMAIL OCULTO]", text)

    # Mask phone numbers (various formats)
    phone_patterns = [
        r"\+?\d{1,3}[-.\s]?\(?\d{1,4}\)?[-.\s]?\d{1,4}[-.\s]?\d{1,9}",
        r"\(\d{2,4}\)\s?\d{4,5}[-.\s]?\d{4}",
        r"\d{10,}",
    ]
    for pattern in phone_patterns:
        text = re.sub(pattern, "[TELÉFONO OCULTO]", text)

    return text


async def parse_cv_with_llm(raw_text: str) -> dict[str, Any]:
    """Parse CV text using LLM to extract structured data."""
    system_prompt = """Eres un experto en análisis de CVs/currículums. Tu tarea es extraer información estructurada del texto del CV proporcionado.

Extrae la siguiente información en formato JSON:
- name: nombre completo del candidato
- email: correo electrónico
- phone: número de teléfono
- location: ubicación/ciudad
- headline: título profesional o resumen en una línea
- summary: resumen profesional (2-3 oraciones)
- skills: lista de habilidades técnicas y blandas
- experience: lista de experiencias laborales con title, company, start_date, end_date, description
- education: lista de educación con degree, institution, year
- languages: lista de idiomas con language y level
- certifications: lista de certificaciones

Si algún campo no está presente en el CV, usa null o lista vacía según corresponda.
Responde SOLO con el JSON, sin explicaciones adicionales."""

    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": f"Analiza el siguiente CV:\n\n{raw_text[:8000]}"},
    ]

    return await llm_provider.complete_json(messages)


def mask_phone(phone: Optional[str]) -> Optional[str]:
    """Mask phone number for display."""
    if not phone:
        return None
    # Show only last 4 digits
    digits = re.sub(r"\D", "", phone)
    if len(digits) >= 4:
        return f"***-***-{digits[-4:]}"
    return "***-***-****"
